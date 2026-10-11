"""Phase 4 acceptance: upload → parse → chunk → embed → index, plus the
failure modes (bad type, oversize, scanned PDF, dup upload, permissions,
idempotent re-ingest)."""

import io
import uuid

import pytest
from fastapi.testclient import TestClient

from tests.test_auth import auth_headers, login, register

PASSWORD = "Sup3rSecret!"


def api_user(client: TestClient, prefix: str) -> dict[str, str]:
    email = f"{prefix}-{uuid.uuid4().hex[:8]}@example.com"
    register(client, email)
    return auth_headers(login(client, email))


def create_doc(client: TestClient, headers: dict[str, str], title: str = "Doc") -> str:
    resp = client.post("/api/v1/documents", headers=headers, json={"title": title})
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def upload(
    client: TestClient,
    headers: dict[str, str],
    doc_id: str,
    data: bytes,
    filename: str,
    effective_from: str = "2024-01-01",
    effective_until: str | None = "2024-12-31",
):
    files = {"file": (filename, io.BytesIO(data))}
    form = {"effective_from": effective_from}
    if effective_until:
        form["effective_until"] = effective_until
    return client.post(
        f"/api/v1/documents/{doc_id}/versions",
        headers=headers,
        files=files,
        data=form,
    )


def make_pdf(text: str = "Annual leave policy. Employees get 12 days.") -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.multi_cell(0, 10, text)
    return bytes(pdf.output())


def make_scanned_pdf() -> bytes:
    """A structurally valid PDF with a blank page — no extractable text."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Leave Policy", level=1)
    d.add_paragraph("Employees are entitled to 12 days of annual leave.")
    d.add_heading("Sick Leave", level=1)
    d.add_paragraph("Employees receive 10 days of sick leave per year.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_txt_upload_ingests_and_chunks(client: TestClient) -> None:
    headers = api_user(client, "ing")
    doc_id = create_doc(client, headers)
    resp = upload(
        client, headers, doc_id,
        b"Paragraph one.\n\nParagraph two about leave.",
        "policy.txt",
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["ingestion_status"] == "completed"
    assert body["status"] == "pending_review"
    assert body["version_number"] == 1
    assert body["extra"]["chunk_count"] >= 1

    chunks = client.get(f"/api/v1/versions/{body['id']}/chunks", headers=headers).json()
    assert len(chunks) == body["extra"]["chunk_count"]
    assert chunks[0]["content"]


def test_pdf_upload_preserves_page_numbers(client: TestClient) -> None:
    headers = api_user(client, "pdf")
    doc_id = create_doc(client, headers)
    resp = upload(client, headers, doc_id, make_pdf(), "policy.pdf")
    assert resp.status_code == 201, resp.text
    version = resp.json()
    assert version["ingestion_status"] == "completed"
    chunks = client.get(
        f"/api/v1/versions/{version['id']}/chunks", headers=headers
    ).json()
    assert all(c["page_number"] == 1 for c in chunks)


def test_docx_upload_preserves_sections(client: TestClient) -> None:
    headers = api_user(client, "docx")
    doc_id = create_doc(client, headers)
    resp = upload(client, headers, doc_id, make_docx(), "policy.docx")
    assert resp.status_code == 201, resp.text
    chunks = client.get(
        f"/api/v1/versions/{resp.json()['id']}/chunks", headers=headers
    ).json()
    sections = {c["section_title"] for c in chunks}
    assert {"Leave Policy", "Sick Leave"} & sections


def test_scanned_pdf_reports_ocr_not_silence(client: TestClient) -> None:
    headers = api_user(client, "scan")
    doc_id = create_doc(client, headers)
    resp = upload(client, headers, doc_id, make_scanned_pdf(), "scan.pdf")
    assert resp.status_code == 201  # version row exists…
    version = resp.json()
    assert version["ingestion_status"] == "failed"  # …but honestly failed
    assert "OCR" in version["extra"]["ingestion_error"]


def test_unsupported_extension_rejected(client: TestClient) -> None:
    headers = api_user(client, "exe")
    doc_id = create_doc(client, headers)
    resp = upload(client, headers, doc_id, b"MZ fake exe", "evil.exe")
    assert resp.status_code == 400
    assert "Supported" in resp.json()["error"]["message"]


def test_extension_lying_about_content_rejected(client: TestClient) -> None:
    headers = api_user(client, "lie")
    doc_id = create_doc(client, headers)
    resp = upload(client, headers, doc_id, b"not a pdf at all", "fake.pdf")
    assert resp.status_code == 400


def test_invalid_interval_rejected(client: TestClient) -> None:
    headers = api_user(client, "date")
    doc_id = create_doc(client, headers)
    resp = upload(
        client, headers, doc_id, b"text",
        "a.txt", effective_from="2025-06-01", effective_until="2024-01-01",
    )
    assert resp.status_code == 400


def test_identical_upload_deduped(client: TestClient) -> None:
    headers = api_user(client, "dupe")
    doc_id = create_doc(client, headers)
    first = upload(client, headers, doc_id, b"same bytes", "a.txt")
    assert first.status_code == 201
    second = upload(client, headers, doc_id, b"same bytes", "a.txt")
    assert second.status_code == 409
    assert "version_number" in second.json()["error"]["details"]


def test_version_numbers_increment(client: TestClient) -> None:
    headers = api_user(client, "inc")
    doc_id = create_doc(client, headers)
    v1 = upload(client, headers, doc_id, b"version one", "a.txt").json()
    v2 = upload(client, headers, doc_id, b"version two", "a.txt").json()
    assert (v1["version_number"], v2["version_number"]) == (1, 2)

    listed = client.get(
        f"/api/v1/documents/{doc_id}/versions", headers=headers
    ).json()
    assert [v["version_number"] for v in listed] == [1, 2]


def test_other_user_cannot_upload_version(client: TestClient) -> None:
    owner_h = api_user(client, "own")
    other_h = api_user(client, "oth")
    doc_id = create_doc(client, owner_h)
    resp = upload(client, other_h, doc_id, b"nope", "a.txt")
    assert resp.status_code == 404  # invisible → not found


def test_reingest_is_idempotent(client: TestClient) -> None:
    headers = api_user(client, "redo")
    doc_id = create_doc(client, headers)
    version = upload(
        client, headers, doc_id, b"chunky content here", "a.txt"
    ).json()
    first_chunks = client.get(
        f"/api/v1/versions/{version['id']}/chunks", headers=headers
    ).json()

    redo = client.post(
        f"/api/v1/versions/{version['id']}/reingest", headers=headers
    )
    assert redo.status_code == 200
    assert redo.json()["ingestion_status"] == "completed"
    second_chunks = client.get(
        f"/api/v1/versions/{version['id']}/chunks", headers=headers
    ).json()
    assert len(second_chunks) == len(first_chunks)


def test_oversize_rejected(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    """25MB of zeros would be silly — shrink the limit instead."""
    from app.core.config import get_settings

    get_settings.cache_clear()
    monkeypatch.setenv("MAX_UPLOAD_MB", "1")  # 1MB for this test
    get_settings.cache_clear()
    try:
        headers = api_user(client, "big")
        doc_id = create_doc(client, headers)
        big = b"x" * (1024 * 1024 + 10)  # just over 1MB
        resp = upload(client, headers, doc_id, big, "big.txt")
        assert resp.status_code == 400
        assert "limit" in resp.json()["error"]["message"].lower()
    finally:
        monkeypatch.undo()
        get_settings.cache_clear()


def test_version_detail_requires_access(client: TestClient) -> None:
    owner_h = api_user(client, "vd1")
    other_h = api_user(client, "vd2")
    doc_id = create_doc(client, owner_h)
    version = upload(client, owner_h, doc_id, b"secret", "a.txt").json()
    resp = client.get(f"/api/v1/versions/{version['id']}", headers=other_h)
    assert resp.status_code == 404
